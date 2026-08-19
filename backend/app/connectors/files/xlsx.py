"""
xlsx — streaming reader for Excel workbooks via openpyxl's read_only
mode, which keeps only the current row's cells resident in memory
instead of materializing the whole sheet (openpyxl's own documented
"optimized reader" for large files). Uses the active sheet's first row
as the header.

XLSX is a ZIP container, which needs random access to its central
directory — a pure network stream can't be forward-only-streamed the
way a CSV can. This still stays off local disk: paramiko's SFTPFile
supports seek()/tell(), so openpyxl can read it directly over SFTP
without this codebase ever downloading the whole workbook first.
"""
from __future__ import annotations

from typing import Iterator

from openpyxl import load_workbook


def peek_header(stream) -> list[str]:
    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = wb.active
        first_row = next(ws.iter_rows(max_row=1, values_only=True), ())
        return [str(v) if v is not None else "" for v in first_row]
    finally:
        wb.close()


def iter_rows(stream) -> Iterator[dict]:
    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            return
        header = [str(v) if v is not None else f"column_{i + 1}" for i, v in enumerate(header_row)]
        for row in rows:
            if row is None or all(v is None for v in row):
                continue
            yield dict(zip(header, row))
    finally:
        wb.close()
